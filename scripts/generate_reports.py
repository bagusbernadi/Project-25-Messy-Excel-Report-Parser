"""Project 25: laporan Excel berantakan (simulasi) dibuat dari data rapi, pabrik minuman kemasan botol. Seed tetap.

Data rapi (truth, 27.000 baris = 150 hari x 2 shift x 3 slot x 10 stasiun x 3 ukuran) ditulis ke 36 file Excel:
6 kelompok stasiun (6 template layout berbeda) x 6 bulan (blok 25 hari kerja). Fakta tertanam: 6 jenis layout,
perubahan layout di bulan ke-4 (layout 1 dan 3), 2 subtotal salah ketik (layout 2), 1 sheet ganda (layout 1 bulan 3),
dan +-2,5% sel angka kotor. Hasil (data/sim/): reports/*.xlsx, truth_fact_output_tidy.csv, ground_truth_dirty_cells.csv,
ground_truth.json. Jalankan: python scripts/generate_reports.py   (versi data: v1; butuh openpyxl)
"""
import io
import json
import re
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "sim"
SEED = 25260302
DATA_VERSION = "v1"
STATIONS = ["Water Treatment", "Mixing", "Blow Molding", "Rinsing", "Filling", "Capping", "Labeling", "Shrink Wrap", "Cartoning", "Palletizing"]
GROUPS = [[0, 1], [2, 3], [4, 5], [6, 7], [8], [9]]               # 6 kelompok -> 6 template layout
SIZES, SHIFTS, SLOTS = ["330 ml", "600 ml", "1,5 L"], ["Shift 1", "Shift 2"], [1, 2, 3]
SIZE_BASE = {"330 ml": 3200, "600 ml": 2800, "1,5 L": 1500}
REJ_RATE = [0.004, 0.006, 0.020, 0.008, 0.015, 0.010, 0.012, 0.005, 0.007, 0.003]
N_BLOCKS, BLOCK_DAYS = 6, 25
LAYOUT_CHANGE_BLOCK = 3                                             # indeks 3 = bulan ke-4
DIRTY_RATE = 0.025
rng = np.random.default_rng(SEED)
days = pd.bdate_range("2026-03-02", periods=N_BLOCKS * BLOCK_DAYS)


def normalize_zip(path):
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in sorted(zin.namelist()):
            zi = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED
            data = zin.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            zout.writestr(zi, data)
    Path(path).write_bytes(buf.getvalue())


# ---- data rapi ----
rows = []
for d in days:
    for sh in SHIFTS:
        for sl in SLOTS:
            for s, st in enumerate(STATIONS):
                for z in SIZES:
                    q = int(round(SIZE_BASE[z] * rng.normal(1, 0.07)))
                    rej = int(rng.binomial(q, REJ_RATE[s]))
                    rows.append([d.date(), sh, sl, st, z, q, rej])
T = pd.DataFrame(rows, columns=["date", "shift", "slot", "station", "size", "qty_produced", "qty_reject"])
T["block"] = [(pd.Timestamp(d) - days[0]).days for d in T["date"]]
T["block"] = [int(np.searchsorted([x.date() for x in days], d)) // BLOCK_DAYS for d in T["date"]]


class Sheet:
    """Pembungkus worksheet yang mencatat sel angka (untuk disisipi data kotor)."""
    def __init__(self, ws):
        self.ws, self.cells = ws, []

    def num(self, r, c, v):
        self.ws.cell(r, c, v); self.cells.append((r, c, v))

    def txt(self, r, c, v):
        self.ws.cell(r, c, v)


def part(block, g):
    d = T[(T.block == block) & (T.station.isin([STATIONS[i] for i in GROUPS[g]]))]
    return d, [x.date() for x in days[block * BLOCK_DAYS:(block + 1) * BLOCK_DAYS]]


def key(d, **kw):
    m = np.ones(len(d), bool)
    for k, v in kw.items(): m &= (d[k] == v).to_numpy()
    return d[m]


def layout1(wb, block, g, info):     # tanggal sebagai kolom; blok PRODUKSI lalu REJECT; judul bergeser di bulan ke-4
    d, ds = part(block, g); ws = wb.active; ws.title = "Sheet1"; sh = Sheet(ws)
    ext = block >= LAYOUT_CHANGE_BLOCK
    ws["A1"] = "LAPORAN PRODUKSI HARIAN"; ws.merge_cells("A1:H1")
    ws["A2"] = f"Area: {', '.join(STATIONS[i] for i in GROUPS[g])}   Periode: {ds[0]:%b %Y}"
    if ext: ws["A3"] = "Revisi form 2026 - mohon isi kolom Catatan"; ws["A4"] = "Dibuat oleh: Supervisor"
    hr = 7 if ext else 5; off = 6 if ext else 4
    r = hr
    for title, col in (("PRODUKSI", "qty_produced"), ("REJECT", "qty_reject")):
        ws.cell(r, 1, title)
        for j, c in enumerate(["Stasiun", "Ukuran", "Shift", "Slot"] + (["Catatan", "Paraf"] if ext else [])):
            ws.cell(r + 1, 1 + j, c)
        for j, x in enumerate(ds): ws.cell(r + 1, off + 1 + j, x.strftime("%d/%m"))
        rr = r + 2; first = rr
        for i in GROUPS[g]:
            for z in SIZES:
                for s in SHIFTS:
                    for sl in SLOTS:
                        for j, c in enumerate([STATIONS[i], z, s, sl]): ws.cell(rr, 1 + j, c)
                        for k, x in enumerate(ds):
                            sh.num(rr, off + 1 + k, int(key(d, date=x, station=STATIONS[i], size=z, shift=s, slot=sl)[col].iloc[0]))
                        rr += 1
        ws.cell(rr, 1, f"TOTAL {title}")
        for k, x in enumerate(ds): ws.cell(rr, off + 1 + k, int(d[d.date == x][col].sum()))
        r = rr + 3
    info["layout_note"] = "tanggal sebagai kolom; blok PRODUKSI dan REJECT; " + ("judul 4 baris + kolom Catatan/Paraf (bulan ke-4)" if ext else "judul 2 baris")
    return [sh], False


def layout2(wb, block, g, info):     # blok per tanggal dan shift, slot melebar, subtotal per shift
    d, ds = part(block, g); ws = wb.active; ws.title = "Laporan"; sh = Sheet(ws)
    r, subtotal_cells = 1, []
    for x in ds:
        ws.cell(r, 1, f"Tanggal: {x:%d %b %Y}"); r += 1
        for s in SHIFTS:
            ws.cell(r, 1, s); r += 1
            for j, c in enumerate(["Stasiun", "Ukuran"] + [f"Slot {sl} {k}" for sl in SLOTS for k in ("Prod", "Rej")] + ["Total Prod"]): ws.cell(r, 1 + j, c)
            r += 1; first = r
            for i in GROUPS[g]:
                for z in SIZES:
                    ws.cell(r, 1, STATIONS[i]); ws.cell(r, 2, z); tot = 0
                    for sl in SLOTS:
                        row = key(d, date=x, station=STATIONS[i], size=z, shift=s, slot=sl).iloc[0]
                        sh.num(r, 3 + (sl - 1) * 2, int(row.qty_produced)); sh.num(r, 4 + (sl - 1) * 2, int(row.qty_reject)); tot += int(row.qty_produced)
                    ws.cell(r, 9, tot); r += 1
            ws.cell(r, 1, f"Subtotal {s}")
            for j, (col, sl) in enumerate([(c, sl) for sl in SLOTS for c in ("qty_produced", "qty_reject")]):
                v = int(key(d, date=x, shift=s, slot=sl)[col].sum()); ws.cell(r, 3 + j, v); subtotal_cells.append((r, 3 + j, v))
            ws.cell(r, 9, int(key(d, date=x, shift=s).qty_produced.sum())); r += 2
    info["layout_note"] = "blok per tanggal dan shift; slot melebar; subtotal per shift (nilai)"
    info["subtotal_cells"] = subtotal_cells
    return [sh], True


def layout3(wb, block, g, info):     # satu sheet per minggu; header dua tingkat digabung
    d, ds = part(block, g); sheets = []
    ext = block >= LAYOUT_CHANGE_BLOCK
    cols = [(STATIONS[i], z) for i in GROUPS[g] for z in SIZES]
    for w in range(5):
        wk = ds[w * 5:(w + 1) * 5]
        ws = wb.active if w == 0 else wb.create_sheet()
        ws.title = f"Wk{w + 1}" if ext else f"Minggu {w + 1}"; sh = Sheet(ws); sheets.append(sh)
        ws["A1"] = f"Output mingguan {wk[0]:%d/%m} - {wk[-1]:%d/%m/%Y}"; ws.merge_cells("A1:F1")
        for j, (st, z) in enumerate(cols):
            ws.cell(3, 4 + j * 2, f"{st} - {z}"); ws.merge_cells(start_row=3, start_column=4 + j * 2, end_row=3, end_column=5 + j * 2)
            ws.cell(4, 4 + j * 2, "Prod"); ws.cell(4, 5 + j * 2, "Rej")
        for j, c in enumerate(["Tanggal", "Shift", "Slot"]): ws.cell(4, 1 + j, c)
        r = 5
        for x in wk:
            for s in SHIFTS:
                for sl in SLOTS:
                    ws.cell(r, 1, x.strftime("%d/%m/%Y")); ws.cell(r, 2, s); ws.cell(r, 3, sl)
                    for j, (st, z) in enumerate(cols):
                        row = key(d, date=x, station=st, size=z, shift=s, slot=sl).iloc[0]
                        sh.num(r, 4 + j * 2, int(row.qty_produced)); sh.num(r, 5 + j * 2, int(row.qty_reject))
                    r += 1
    info["layout_note"] = "satu sheet per minggu; header dua tingkat digabung" + ("; nama sheet 'Wk1' (bulan ke-4)" if ext else "; nama sheet 'Minggu 1'")
    return sheets, False


def layout4(wb, block, g, info):     # Shift 1 di kiri, Shift 2 di kanan, kolom pemisah kosong
    d, ds = part(block, g); ws = wb.active; ws.title = "Output"; sh = Sheet(ws)
    for off, s in ((0, "Shift 1"), (7, "Shift 2")):
        ws.cell(1, 1 + off, s)
        for j, c in enumerate(["Tanggal", "Slot", "Stasiun", "Ukuran", "Produksi", "Reject"]): ws.cell(2, 1 + off + j, c)
        r = 3
        for x in ds:
            for sl in SLOTS:
                for i in GROUPS[g]:
                    for z in SIZES:
                        row = key(d, date=x, station=STATIONS[i], size=z, shift=s, slot=sl).iloc[0]
                        ws.cell(r, 1 + off, x.strftime("%d/%m/%Y")); ws.cell(r, 2 + off, sl); ws.cell(r, 3 + off, STATIONS[i]); ws.cell(r, 4 + off, z)
                        sh.num(r, 5 + off, int(row.qty_produced)); sh.num(r, 6 + off, int(row.qty_reject)); r += 1
    info["layout_note"] = "Shift 1 di kolom A-F, Shift 2 di kolom H-M, kolom G pemisah"
    return [sh], False


ALIAS = {"Water Treatment": ["WTP", "Water treatment"], "Mixing": ["Mixing", "Mixer"], "Blow Molding": ["Blow", "Blow Molding"], "Rinsing": ["Rinser", "Rinsing"],
         "Filling": ["Filler", "Filling"], "Capping": ["Capper", "Capping"], "Labeling": ["Label", "Labeling"], "Shrink Wrap": ["Shrink", "Shrink Wrap"],
         "Cartoning": ["Carton", "Cartoning"], "Palletizing": ["Palet", "Palletizing"]}
SIZE_ALIAS = {"330 ml": ["330ml", "330 ML", "0,33 L"], "600 ml": ["600ml", "600 ML", "0,6 L"], "1,5 L": ["1.5L", "1500 ml", "1,5 Liter"]}


def layout5(wb, block, g, info):     # tabel rata dengan angka sebagai teks dan penulisan tidak seragam
    d, ds = part(block, g); ws = wb.active; ws.title = "Data"; sh = Sheet(ws)
    for j, c in enumerate(["Tgl", "Shift", "Jam", "Mesin", "Botol", "Hasil", "Afkir"]): ws.cell(1, 1 + j, c)
    r = 2; n = 0
    for x in ds:
        for s in SHIFTS:
            for sl in SLOTS:
                for i in GROUPS[g]:
                    for z in SIZES:
                        row = key(d, date=x, station=STATIONS[i], size=z, shift=s, slot=sl).iloc[0]
                        fmt = [x.strftime("%d-%b-%y"), x.strftime("%Y/%m/%d"), x.strftime("%d.%m.%Y")][n % 3]
                        ws.cell(r, 1, fmt); ws.cell(r, 2, [s, s[-1], "I" if s == "Shift 1" else "II"][n % 3])
                        ws.cell(r, 3, ["07:00-10:00", "10:00-13:00", "13:00-16:00"][sl - 1] if s == "Shift 1" else ["15:00-18:00", "18:00-21:00", "21:00-24:00"][sl - 1])
                        ws.cell(r, 4, ALIAS[STATIONS[i]][n % 2]); ws.cell(r, 5, SIZE_ALIAS[z][n % 3])
                        ws.cell(r, 6, f"{int(row.qty_produced):,}".replace(",", ".") + " btl"); ws.cell(r, 7, f"{int(row.qty_reject)} pcs")
                        r += 1; n += 1
    info["layout_note"] = "tabel rata; tanggal/shift/stasiun/ukuran ditulis tidak seragam; angka sebagai teks dengan satuan"
    return [], False


def layout6(wb, block, g, info):     # stasiun melebar (header digabung), ukuran sebagai baris, baris bantu tersembunyi, catatan kaki
    d, ds = part(block, g); ws = wb.active; ws.title = "Rekap"; sh = Sheet(ws)
    sts = [STATIONS[i] for i in GROUPS[g]]
    ws["A1"] = "REKAP OUTPUT"; ws.append([])
    ws.row_dimensions[2].hidden = True
    for j, st in enumerate(sts):
        ws.cell(2, 5 + j * 2, f"S{GROUPS[g][j]:02d}"); ws.cell(3, 5 + j * 2, st); ws.merge_cells(start_row=3, start_column=5 + j * 2, end_row=3, end_column=6 + j * 2)
        ws.cell(4, 5 + j * 2, "Prod"); ws.cell(4, 6 + j * 2, "Rej")
    for j, c in enumerate(["Tanggal", "Shift", "Slot", "Ukuran"]): ws.cell(4, 1 + j, c)
    r = 5
    for x in ds:
        for s in SHIFTS:
            for sl in SLOTS:
                for z in SIZES:
                    ws.cell(r, 1, x); ws.cell(r, 2, s); ws.cell(r, 3, sl); ws.cell(r, 4, z)
                    ws.cell(r, 1).number_format = "dd/mm/yyyy"
                    for j, st in enumerate(sts):
                        row = key(d, date=x, station=st, size=z, shift=s, slot=sl).iloc[0]
                        sh.num(r, 5 + j * 2, int(row.qty_produced)); sh.num(r, 6 + j * 2, int(row.qty_reject))
                    r += 1
    ws.cell(r + 1, 1, "Catatan: angka dalam botol. Reject = afkir sebelum QC akhir."); ws.cell(r + 2, 1, "Disetujui: ____________")
    info["layout_note"] = "stasiun melebar (header digabung); ukuran sebagai baris; baris 2 tersembunyi; catatan kaki"
    return [sh], False


LAYOUTS = [layout1, layout2, layout3, layout4, layout5, layout6]
(OUT / "reports").mkdir(parents=True, exist_ok=True)
files, dirty_log = [], []
sub_typos_left = [1, 4]                                          # blok bulan (indeks) layout 2 yang subtotalnya salah ketik
for b in range(N_BLOCKS):
    for g in range(6):
        info = {"group": g + 1, "stations": [STATIONS[i] for i in GROUPS[g]], "block": b + 1, "layout": g + 1}
        wb = Workbook(); wb.properties.creator = "simulasi"
        sheets, has_sub = LAYOUTS[g](wb, b, g, info)
        fname = f"laporan_G{g + 1}_{days[b * BLOCK_DAYS]:%Y-%m}.xlsx"
        if has_sub and b in sub_typos_left:                      # subtotal salah ketik: dua digit tertukar
            r, c, v = info["subtotal_cells"][int(rng.integers(0, len(info["subtotal_cells"])))]
            s = str(v); bad = int(s[0] + s[2] + s[1] + s[3:]) if len(s) >= 4 and s[1] != s[2] else v + 90
            wb.worksheets[0].cell(r, c, bad)
            dirty_log.append([fname, wb.worksheets[0].title, f"{L(c)}{r}", "subtotal_typo", v, bad])
        info.pop("subtotal_cells", None)
        if g == 0 and b == 2:                                    # sheet ganda
            copy = wb.copy_worksheet(wb.worksheets[0]); copy.title = "Sheet1 (2)"
            dirty_log.append([fname, "Sheet1 (2)", "", "duplicate_sheet", "", ""])
        allcells = [(sh.ws.title, r, c, v) for sh in sheets for (r, c, v) in sh.cells]
        if allcells:
            pick = rng.permutation(len(allcells))[: int(round(len(allcells) * DIRTY_RATE))]
            for k, i in enumerate(pick):
                title, r, c, v = allcells[int(i)]
                kind = ["text_na", "dash", "blank", "x10", "unit_text"][k % 5]
                new = {"text_na": "n/a", "dash": "-", "blank": None, "x10": v * 10, "unit_text": f"{v} btl"}[kind]
                wb[title].cell(r, c, new)
                dirty_log.append([fname, title, f"{L(c)}{r}", kind, v, new])
        wb.save(OUT / "reports" / fname); normalize_zip(OUT / "reports" / fname)
        files.append({"file": fname, **info, "sheets": [w.title for w in wb.worksheets]})

T.drop(columns="block").to_csv(OUT / "truth_fact_output_tidy.csv", index=False)
pd.DataFrame(dirty_log, columns=["file", "sheet", "cell", "problem", "original", "written"]).to_csv(OUT / "ground_truth_dirty_cells.csv", index=False)
(OUT / "ground_truth.json").write_text(json.dumps({
    "seed": SEED, "version": DATA_VERSION, "tidy_rows": len(T), "n_files": len(files), "n_problem_cells": len(dirty_log),
    "facts": [{"id": 1, "fact": "6 template layout berbeda, satu per kelompok stasiun"},
              {"id": 2, "fact": "Layout 1 dan 3 berubah di bulan ke-4", "from": str(days[LAYOUT_CHANGE_BLOCK * BLOCK_DAYS].date())},
              {"id": 3, "fact": "2 subtotal salah ketik (layout 2)", "blocks": [b + 1 for b in sub_typos_left]},
              {"id": 4, "fact": "1 sheet ganda", "file": "laporan_G1_2026-05.xlsx"},
              {"id": 5, "fact": "Sel angka kotor", "rate": DIRTY_RATE}],
    "files": files}, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
print(f"{len(files)} file, tidy {len(T)} baris, sel bermasalah {len(dirty_log)}")
