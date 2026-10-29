"""Tahap 1 / langkah 2: memeriksa 36 laporan apa adanya (tanpa parser) dan mendata masalah layout (Q1).

Hasil: results/1-data-preparation/{layout_features.csv, layout_taxonomy.csv, data_profile.csv}
"""
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import common as C  # noqa: E402

DATE_HDR = re.compile(r"^\d{2}/\d{2}$")
NUMTEXT = re.compile(r"^\s*[\d.,]+\s*[A-Za-z]*\s*$")
PROBLEMS = [("merged_header", "Header digabung (merged cell) di dalam tabel"),
            ("dates_as_columns", "Tanggal berjajar sebagai kolom"),
            ("multi_table", "Lebih dari satu tabel dalam satu sheet"),
            ("subtotal_inside", "Baris subtotal atau total di tengah data"),
            ("sheet_per_period", "Satu sheet per minggu"),
            ("text_numbers", "Angka tersimpan sebagai teks dengan satuan"),
            ("inconsistent_labels", "Nama stasiun, ukuran, atau tanggal ditulis tidak seragam"),
            ("hidden_or_footnote", "Baris tersembunyi atau catatan kaki"),
            ("extra_columns", "Kolom catatan atau paraf di luar data")]


def inspect(path):
    wb = openpyxl.load_workbook(path)
    f = dict.fromkeys([p for p, _ in PROBLEMS], False)
    n_cells = n_num = n_numtext = n_merged = n_hidden = n_hdr_merged = 0
    sheet_names, labels, datefmts = wb.sheetnames, set(), set()
    for ws in wb.worksheets:
        n_merged += len(ws.merged_cells.ranges)
        n_hdr_merged += sum(1 for m in ws.merged_cells.ranges if m.min_row > 1)      # judul di baris 1 tidak dihitung
        n_hidden += sum(1 for d in ws.row_dimensions.values() if d.hidden)
        rows_a = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                n_cells += 1
                if isinstance(v, (int, float)):
                    n_num += 1
                elif isinstance(v, str):
                    if NUMTEXT.match(v) and re.search(r"\d", v):
                        n_numtext += 1
                    if c.column == 1:
                        rows_a[c.row] = v
                    if DATE_HDR.match(v):
                        f["dates_as_columns"] = True
                    if re.match(r"^\d{1,2}[-/.]\w{2,3}[-/.]\d{2,4}$|^\d{4}/\d{2}/\d{2}$", v):
                        datefmts.add(re.sub(r"[A-Za-z]+", "a", re.sub(r"\d", "9", v)))
        vals_a = list(rows_a.values())
        if any(v.startswith(("TOTAL", "Subtotal")) for v in vals_a):
            f["subtotal_inside"] = True
        if sum(v in ("PRODUKSI", "REJECT") for v in vals_a) >= 2 or (ws["A1"].value == "Shift 1" and ws["H1"].value == "Shift 2"):
            f["multi_table"] = True
        if any(v.startswith(("Catatan:", "Disetujui")) for v in vals_a):
            f["hidden_or_footnote"] = True
        for c in ws[ws.min_row]:
            if c.value in ("Catatan", "Paraf"):
                f["extra_columns"] = True
        for row in ws.iter_rows(min_row=1, max_row=12):
            for c in row:
                if c.value in ("Catatan", "Paraf"):
                    f["extra_columns"] = True
        for col in ("D", "E") if ws.title == "Data" else ():
            labels |= {str(ws[f"{col}{r}"].value) for r in range(2, ws.max_row + 1)}
    f["merged_header"] = n_hdr_merged > 0
    f["sheet_per_period"] = len(sheet_names) >= 4
    f["text_numbers"] = n_numtext / max(n_numtext + n_num, 1) > 0.5
    f["inconsistent_labels"] = len(labels) > 6 or len(datefmts) > 1
    f["hidden_or_footnote"] = f["hidden_or_footnote"] or n_hidden > 0
    return {"n_sheets": len(sheet_names), "n_cells_with_value": n_cells, "n_numeric_cells": n_num, "n_text_number_cells": n_numtext, "n_merged_ranges": n_merged,
            "n_hidden_rows": n_hidden, "n_label_variants": len(labels), "n_date_formats": len(datefmts), **f}


files = sorted(C.REPORTS_DIR.glob("*.xlsx"))
assert len(files) == C.N_FILES, f"{len(files)} file di {C.REPORTS_DIR}; jalankan 01_generate_reports.py"
rows = []
for p in files:
    m = re.match(r"laporan_G(\d)_(\d{4}-\d{2})\.xlsx", p.name)
    rows.append({"file": p.name, "layout": int(m.group(1)), "month": m.group(2), **inspect(p)})
feat = pd.DataFrame(rows)
feat.to_csv(C.RES_PREP / "layout_features.csv", index=False)

tax = []
for key, label in PROBLEMS:
    row = {"problem": key, "description": label, "files": int(feat[key].sum())}
    for lay in range(1, 7):
        row[f"layout{lay}"] = int(feat[feat.layout == lay][key].sum())
    tax.append(row)
pd.DataFrame(tax).sort_values(["files", "problem"], ascending=[False, True], kind="stable").to_csv(C.RES_PREP / "layout_taxonomy.csv", index=False)

prof = [("file laporan", len(feat)), ("sheet", int(feat.n_sheets.sum())), ("sel berisi nilai", int(feat.n_cells_with_value.sum())),
        ("sel angka", int(feat.n_numeric_cells.sum())), ("sel angka bertipe teks", int(feat.n_text_number_cells.sum())), ("merged cell", int(feat.n_merged_ranges.sum())),
        ("baris tersembunyi", int(feat.n_hidden_rows.sum())), ("layout", int(feat.layout.nunique())), ("bulan", int(feat.month.nunique()))]
pd.DataFrame(prof, columns=["item", "value"]).to_csv(C.RES_PREP / "data_profile.csv", index=False)
print(pd.DataFrame(tax).to_string(index=False))
