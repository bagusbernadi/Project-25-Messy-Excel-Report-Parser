"""Memuat config template, kamus nama, dan aturan; mengenali template sebuah workbook (tanpa menebak)."""
import re
from pathlib import Path

import yaml
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils.cell import coordinate_from_string


class ParseError(Exception):
    """Struktur laporan tidak sesuai dengan config."""


class UnknownTemplate(ParseError):
    """Tidak ada config yang cocok dengan workbook ini."""


class AmbiguousTemplate(ParseError):
    """Lebih dari satu config cocok."""


class UnknownLabel(ParseError):
    """Nama stasiun, ukuran, atau shift yang tidak ada di lookup.yaml."""


def _key(s):
    return re.sub(r"\s+", " ", str(s).strip()).casefold()


def cell_rc(addr):
    col, row = coordinate_from_string(addr)
    return row, CI(col)


class Lookup:
    def __init__(self, d):
        self.d = d
        self.stations = {_key(a): c for c, al in d["stations"].items() for a in al}
        self.sizes = {_key(a): c for c, al in d["sizes"].items() for a in al}
        self.shifts = {_key(a): c for c, al in d["shifts"].items() for a in al}
        self.months = {k.casefold(): v for k, v in d["months"].items()}
        self.slot_by_time = d["slot_by_time"]

    def _get(self, table, kind, v):
        try:
            return table[_key(v)]
        except KeyError:
            raise UnknownLabel(f"{kind} tidak dikenal: {v!r} (tambahkan ke config/lookup.yaml)") from None

    def station(self, v):
        return self._get(self.stations, "stasiun", v)

    def size(self, v):
        return self._get(self.sizes, "ukuran", v)

    def shift(self, v):
        return self._get(self.shifts, "shift", v)

    def month(self, v):
        return self._get(self.months, "bulan", v)

    def slot_from_time(self, shift, text):
        try:
            return self.slot_by_time[shift][re.sub(r"\s+", "", str(text))]
        except KeyError:
            raise UnknownLabel(f"jam tidak dikenal: {text!r} untuk {shift}") from None


def load_yaml(p):
    return yaml.safe_load(Path(p).read_text(encoding="utf-8"))


def load_lookup(p):
    return Lookup(load_yaml(p))


def load_rules(p):
    return load_yaml(p)


def load_templates(directory):
    """Semua config template di folder (urut nama file)."""
    out = []
    for p in sorted(Path(directory).glob("*.yaml")):
        cfg = load_yaml(p)
        cfg["_file"] = p.name
        out.append(cfg)
    if not out:
        raise ParseError(f"tidak ada config di {directory}")
    return out


def _cell_ok(ws, rule):
    v = ws[rule["cell"]].value
    s = "" if v is None else str(v).strip()
    if "equals" in rule:
        return s == str(rule["equals"])
    if "startswith" in rule:
        return s.startswith(str(rule["startswith"]))
    if "contains" in rule:
        return str(rule["contains"]) in s
    if rule.get("empty"):
        return s == ""
    raise ParseError(f"aturan detect tidak dikenal: {rule}")


def matches(wb, cfg):
    det = cfg["detect"]
    names = [n for n in wb.sheetnames if re.search(det["sheet_regex"], n)]
    if not names:
        return False
    ws = wb[names[0]]
    return all(_cell_ok(ws, r) for r in det.get("cells", []))


def detect_template(wb, templates):
    hit = [c for c in templates if matches(wb, c)]
    if not hit:
        raise UnknownTemplate(f"tidak ada template yang cocok; sheet: {wb.sheetnames}")
    if len(hit) > 1:
        raise AmbiguousTemplate("lebih dari satu template cocok: " + ", ".join(c["id"] for c in hit))
    return hit[0]
