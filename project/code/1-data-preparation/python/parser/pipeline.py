"""Alur parsing: kenali template, baca sheet, bersihkan angka, tandai outlier, rekonsiliasi, susun tabel hasil."""
import re
import time
from pathlib import Path

import openpyxl
import pandas as pd

from .clean import clean_value, flag_outliers, outlier_stats
from .config import ParseError, detect_template
from .readers import READERS
from .reconcile import reconcile


def _similarity(a, b):
    ca = {(c.row, c.column): c.value for row in a.iter_rows() for c in row if c.value is not None}
    cb = {(c.row, c.column): c.value for row in b.iter_rows() for c in row if c.value is not None}
    keys = set(ca) | set(cb)
    return sum(ca.get(k) == cb.get(k) for k in keys) / len(keys) if keys else 1.0


def parse_workbook(path, templates, lk, rules):
    """Baca satu file: records mentah (sudah dibersihkan angkanya) + constraints + catatan sheet."""
    path = Path(path)
    wb = openpyxl.load_workbook(path)                     # bukan read_only: merged cell dan baris tersembunyi perlu dibaca
    cfg = detect_template(wb, templates)
    t0 = time.perf_counter()
    recs, cons, sheet_notes = [], [], []
    dup = rules["duplicate_sheet"]
    for name in wb.sheetnames:
        if not re.search(cfg["sheets"]["include"], name):
            base = re.sub(dup["name_regex"], "", name)
            if base != name and base in wb.sheetnames:
                sim = _similarity(wb[base], wb[name])
                if sim >= dup["min_similarity"]:
                    sheet_notes.append({"sheet": name, "issue_type": "duplicate_sheet", "raw_value": f"mirip {sim:.1%} dengan '{base}'", "action": "dropped"})
                else:
                    sheet_notes.append({"sheet": name, "issue_type": "duplicate_sheet", "raw_value": f"nama salinan tetapi hanya mirip {sim:.1%}", "action": "flagged"})
            else:
                sheet_notes.append({"sheet": name, "issue_type": "unexpected_sheet", "raw_value": "", "action": "flagged"})
            continue
        off = len(recs)
        r_, c_ = READERS[cfg["reader"]](wb[name], cfg, lk)
        for c in c_:
            c["members"] = [i + off for i in c["members"]]
        recs += r_
        cons += c_
    if not recs:
        raise ParseError(f"{path.name}: tidak ada sheet data yang terbaca dengan {cfg['id']}")
    for r in recs:
        r["file"], r["template"], r["layout"] = path.name, cfg["id"], cfg["layout"]
        r["value"], r["issue"], r["fix"] = clean_value(r["raw"], rules, cfg.get("text_numbers_expected", False))
        r.update(suspect=False, fixed_by=None, cleared=False, z=0.0)
    for c in cons:
        c["file"] = path.name
    return {"file": path.name, "template": cfg["id"], "layout": cfg["layout"], "records": recs, "constraints": cons, "sheet_notes": sheet_notes,
            "seconds_read": time.perf_counter() - t0}


def parse_files(paths, templates, lk, rules):
    """Parse banyak file. Mengembalikan dict berisi DataFrame fact, anomaly, reconciliation, files."""
    parsed = [parse_workbook(p, templates, lk, rules) for p in paths]
    allrecs = [r for p in parsed for r in p["records"]]
    flag_outliers(allrecs, outlier_stats(allrecs), rules)
    fact, anom, rec_rows, files = [], [], [], []
    for p in parsed:
        recs, cons = p["records"], reconcile(p["records"], p["constraints"], rules)
        for r in recs:
            if r["fixed_by"] is not None or (r["fix"] == "strip_unit" and r["value"] is not None):
                q = "fixed"
            elif r["value"] is None:
                q = "missing"
            elif r["suspect"]:
                q = "flagged"
            else:
                q = "ok"
            r["quality"] = q
            fact.append({"date": r["date"], "shift": r["shift"], "slot": r["slot"], "station": r["station"], "size": r["size"], "metric": r["metric"],
                         "value": r["value"], "quality": q, "template": r["template"], "source_file": r["file"], "source_sheet": r["sheet"], "source_cell": r["cell"]})
            if q != "ok":
                if r["fixed_by"] is not None:
                    how = "subtotal"
                elif q == "fixed":
                    how = "strip_unit"
                else:
                    how = ""
                anom.append({"source_file": r["file"], "sheet": r["sheet"], "cell": r["cell"], "template": r["template"],
                             "issue_type": r["issue"] or "outlier", "raw_value": "" if r["raw"] is None else str(r["raw"]),
                             "action": "fixed" if q == "fixed" else "flagged", "fix_method": how,
                             "parsed_value": r["value"]})
        for c in cons:
            rec_rows.append({"source_file": c["file"], "sheet": c["sheet"], "block": c["block"], "metric": c["metric"], "kind": c["kind"], "report_cell": c["report_cell"],
                             "report_total": c["report_total"], "parsed_total": c["parsed_raw"], "diff": c["diff_raw"], "status": c["status_raw"],
                             "parsed_total_final": c["parsed_final"], "diff_final": c["diff_final"], "status_final": c["status_final"], "resolution": c["resolution"]})
            if c["resolution"] == "selisih: subtotal dicurigai":
                anom.append({"source_file": c["file"], "sheet": c["sheet"], "cell": c["report_cell"], "template": p["template"], "issue_type": "subtotal_mismatch",
                             "raw_value": str(c["report_raw"]), "action": "flagged", "fix_method": "", "parsed_value": c["parsed_final"]})
        for n in p["sheet_notes"]:
            anom.append({"source_file": p["file"], "sheet": n["sheet"], "cell": "", "template": p["template"], "issue_type": n["issue_type"], "raw_value": n["raw_value"],
                         "action": n["action"], "fix_method": "", "parsed_value": None})
        files.append({"source_file": p["file"], "template": p["template"], "layout": p["layout"], "n_cells": len(recs), "n_constraints": len(cons)})
    return {"fact": pd.DataFrame(fact), "anomaly": pd.DataFrame(anom), "reconciliation": pd.DataFrame(rec_rows), "files": pd.DataFrame(files),
            "seconds_read": {p["file"]: p["seconds_read"] for p in parsed}}
