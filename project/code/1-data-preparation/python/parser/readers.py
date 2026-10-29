"""Empat jenis pembaca. Config memilih pembaca dan mengisi parameternya (letak sel, penanda, nama kolom).

Tiap pembaca mengembalikan (records, constraints):
- record: satu sel angka dengan kunci (date, shift, slot, station, size, metric), alamat sel, dan isi mentah.
- constraint: satu subtotal laporan (atau total baris) beserta indeks record yang harus berjumlah sama dengannya.
"""
import re
from datetime import date

from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as L

from .cells import Grid, parse_date
from .config import ParseError, cell_rc


def _rec(sheet, r, c, metric, raw, **key):
    return {"sheet": sheet, "row": r, "col": c, "cell": f"{L(c)}{r}", "metric": metric, "raw": raw, **key}


def _con(kind, sheet, block, metric, r, c, raw, members):
    return {"kind": kind, "sheet": sheet, "block": block, "metric": metric, "report_cell": f"{L(c)}{r}", "report_raw": raw, "members": members}


def _norm(d, lk):
    out = dict(d)
    if "station" in out:
        out["station"] = lk.station(out["station"])
    if "size" in out:
        out["size"] = lk.size(out["size"])
    if "shift" in out:
        out["shift"] = lk.shift(out["shift"])
    if "slot" in out:
        out["slot"] = int(out["slot"])
    return out


def read_date_columns(ws, cfg, lk):
    """Tanggal sebagai kolom; beberapa blok (PRODUKSI, REJECT) di satu sheet; baris TOTAL di bawah tiap blok."""
    g, recs, cons = Grid(ws), [], []
    m = re.search(cfg["period"]["regex"], str(g.v(*cell_rc(cfg["period"]["cell"]))))
    if not m:
        raise ParseError(f"periode tidak ditemukan di {cfg['period']['cell']}")
    mon0, year0 = lk.month(m.group(1)), int(m.group(2))
    mcol, v0 = CI(cfg["marker_col"]), CI(cfg["first_value_col"])
    kcols = {CI(k): f for k, f in cfg["key_columns"].items()}
    markers = [(r, str(g.v(r, mcol)).strip()) for r in range(1, g.max_row + 1) if str(g.v(r, mcol)).strip() in cfg["blocks"]]
    if len(markers) != len(cfg["blocks"]):
        raise ParseError(f"penanda blok {list(cfg['blocks'])} tidak lengkap di kolom {cfg['marker_col']}")
    for r0, name in markers:
        metric, hdr = cfg["blocks"][name], r0 + 1
        dcols, c = [], v0
        while g.v(hdr, c) not in (None, ""):
            dcols.append(c)
            c += 1
        dates = []
        for c in dcols:
            dd, mm = (int(x) for x in str(g.v(hdr, c)).split("/"))
            dates.append(date(year0 + (1 if mm < mon0 else 0), mm, dd))
        rows, r, total_row = [], hdr + 1, None
        while r <= g.max_row and g.v(r, mcol) is not None:
            if str(g.v(r, mcol)).startswith(cfg["subtotal_prefix"]):
                total_row = r
                break
            rows.append(r)
            r += 1
        keys = {r: _norm({f: g.v(r, ci) for ci, f in kcols.items()}, lk) for r in rows}
        for c, d in zip(dcols, dates):
            members = []
            for r in rows:
                members.append(len(recs))
                recs.append(_rec(ws.title, r, c, metric, g.v(r, c), date=d, **keys[r]))
            if total_row:
                cons.append(_con("subtotal", ws.title, f"{d.isoformat()}|{name}", metric, total_row, c, g.v(total_row, c), members))
    return recs, cons


def read_date_shift_groups(ws, cfg, lk):
    """Blok berulang per tanggal dan shift; slot melebar; baris Subtotal per shift dan kolom total produksi per baris."""
    g, recs, cons = Grid(ws), [], []
    lc = CI(cfg["label_col"])
    cols = cfg["columns"]
    sc, zc, tc = CI(cols["station"]), CI(cols["size"]), CI(cols["row_total_produced"])
    slots = {int(s): (CI(p), CI(q)) for s, (p, q) in cols["slots"].items()}
    cur_date = cur_shift = None
    block = []                                    # indeks record baris data di blok shift berjalan: [(row, {slot: (ip, iq)})]
    for r in range(1, g.max_row + 1):
        a = str(g.v(r, lc) or "").strip()
        m = re.match(cfg["date_marker"], a)
        if m:
            cur_date, cur_shift, block = date(int(m.group(3)), lk.month(m.group(2)), int(m.group(1))), None, []
            continue
        if re.match(cfg["shift_marker"], a):
            cur_shift, block = lk.shift(a), []
            continue
        if not a or a == cfg["header_label"]:
            continue
        if a.startswith(cfg["subtotal_prefix"]):
            blk = f"{cur_date.isoformat()}|{cur_shift}"
            for s, (pc, qc) in slots.items():
                for metric, col, ix in (("produced", pc, 0), ("reject", qc, 1)):
                    cons.append(_con("subtotal", ws.title, f"{blk}|slot {s}", metric, r, col, g.v(r, col), [row[1][s][ix] for row in block]))
            cons.append(_con("subtotal", ws.title, f"{blk}|total produksi", "produced", r, tc, g.v(r, tc), [row[1][s][0] for row in block for s in slots]))
            block = []
            continue
        if cur_date is None:
            continue                                      # judul di atas blok pertama
        if cur_shift is None:
            raise ParseError(f"baris data di luar blok shift (baris {r})")
        key = _norm({"station": g.v(r, sc), "size": g.v(r, zc)}, lk)
        idx = {}
        for s, (pc, qc) in slots.items():
            ip = len(recs)
            recs.append(_rec(ws.title, r, pc, "produced", g.v(r, pc), date=cur_date, shift=cur_shift, slot=s, **key))
            recs.append(_rec(ws.title, r, qc, "reject", g.v(r, qc), date=cur_date, shift=cur_shift, slot=s, **key))
            idx[s] = (ip, ip + 1)
        cons.append(_con("row_total", ws.title, f"{cur_date.isoformat()}|{cur_shift}|{key['station']}|{key['size']}", "produced", r, tc, g.v(r, tc), [idx[s][0] for s in slots]))
        block.append((r, idx))
    return recs, cons


def read_wide_groups(ws, cfg, lk):
    """Header dua tingkat (grup digabung + Prod/Rej); tiap baris adalah kombinasi tanggal, shift, slot (dan ukuran)."""
    g, recs = Grid(ws), []
    gr, sr, fr = cfg["group_header_row"], cfg["sub_header_row"], cfg["first_data_row"]
    v0 = CI(cfg["first_value_col"])
    kcols = {CI(k): f for k, f in cfg["key_columns"].items()}
    gh = cfg["group_header"]
    cols = []
    for c in range(v0, g.max_col + 1):
        top, sub = g.v(gr, c, fill=True), g.v(sr, c)
        if top is None or sub is None:
            continue
        parts = str(top).split(gh["split"]) if gh.get("split") else [str(top)]
        if len(parts) != len(gh["fields"]):
            raise ParseError(f"header grup {top!r} tidak sesuai {gh['fields']}")
        if str(sub).strip() not in cfg["sub_headers"]:
            raise ParseError(f"sub-header tidak dikenal: {sub!r}")
        cols.append((c, cfg["sub_headers"][str(sub).strip()], _norm(dict(zip(gh["fields"], (p.strip() for p in parts))), lk)))
    if not cols:
        raise ParseError("tidak ada kolom nilai di bawah header grup")
    for r in range(fr, g.max_row + 1):
        if cfg.get("skip_hidden_rows") and r in g.hidden:
            continue
        kv = {f: g.v(r, ci) for ci, f in kcols.items()}
        d = parse_date(kv["date"], cfg["date_formats"], lk.months)
        if d is None:
            if cfg.get("stop_when_date_unparseable") or all(v is None for v in kv.values()):
                break
            raise ParseError(f"tanggal tidak terbaca di baris {r}: {kv['date']!r}")
        kv["date"] = d
        key = _norm(kv, lk)
        for c, metric, gkey in cols:
            recs.append(_rec(ws.title, r, c, metric, g.v(r, c), **{**key, **gkey}))
    return recs, []


def read_flat(ws, cfg, lk):
    """Tabel rata; satu atau beberapa blok berdampingan; angka boleh berupa teks."""
    g, recs = Grid(ws), []
    names = cfg["columns"]
    fr = cfg["first_data_row"]
    for blk in cfg["blocks"]:
        sc = CI(blk["start_col"])
        title_shift = g.v(*cell_rc(blk["shift_cell"])) if blk.get("shift_cell") else None
        for r in range(fr, g.max_row + 1):
            vals = {n: g.v(r, sc + i) for i, n in enumerate(names)}
            if all(v is None for v in vals.values()):
                break
            d = parse_date(vals["date"], cfg["date_formats"], lk.months)
            if d is None:
                raise ParseError(f"tanggal tidak terbaca di baris {r}: {vals['date']!r}")
            shift = lk.shift(title_shift if title_shift is not None else vals["shift"])
            slot = vals["slot"] if "slot" in vals else lk.slot_from_time(shift, vals["time"])
            key = _norm({"station": vals["station"], "size": vals["size"], "slot": slot}, lk)
            for metric in ("produced", "reject"):
                recs.append(_rec(ws.title, r, sc + names.index(metric), metric, vals[metric], date=d, shift=shift, **key))
    return recs, []


READERS = {"date_columns": read_date_columns, "date_shift_groups": read_date_shift_groups, "wide_groups": read_wide_groups, "flat": read_flat}
