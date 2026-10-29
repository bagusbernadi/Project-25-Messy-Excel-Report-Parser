"""Skenario perubahan layout (Q6, Q7): laporan dimodifikasi, lalu diparse dengan config lama dan config yang diperbarui.

Perubahan config ditulis sebagai penggantian satu baris teks YAML (apa yang akan diedit orang), sehingga jumlah baris yang diubah bisa dihitung.
Dipakai oleh 08_robustness.py dan tests/test_scenarios.py.
"""
import difflib
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analysis_lib import C, KEY, load_truth_long  # noqa: E402
from parser import ParseError, UnknownLabel, UnknownTemplate, load_lookup, load_rules, load_templates, parse_files  # noqa: E402
from parser.config import Lookup  # noqa: E402


def m_insert_col_front(ws_name):
    def f(wb):
        wb[ws_name].insert_cols(1)
    return f


def m_title_rows(wb):
    wb["Laporan"].insert_rows(1, 2)
    wb["Laporan"]["A1"] = "LAPORAN SHIFT PABRIK"


def m_insert_col_d(wb):
    wb["Data"].insert_cols(4)
    wb["Data"]["D1"] = "Operator"
    for r in range(2, wb["Data"].max_row + 1):
        wb["Data"].cell(r, 4, "-")


def m_rename_sheet(wb):
    wb["Sheet1"].title = "Output"


def m_extra_sheet(wb):
    ws = wb.create_sheet("Catatan")
    ws["A1"] = "Catatan supervisor bulan ini"


def m_unknown_station(wb):
    for ws in wb.worksheets:
        for c in ws[3]:
            if isinstance(c.value, str) and c.value.startswith("Filling"):
                c.value = c.value.replace("Filling", "Filler 2")


def m_random(wb):
    for ws in list(wb.worksheets):
        wb.remove(ws)
    ws = wb.create_sheet("Hitungan")
    ws["A1"], ws["B1"] = "Bukan laporan produksi", 3


def _lines(text):
    return [ln for ln in text.splitlines() if ln.strip() and not ln.startswith(("id:", "description:"))]


SCEN = [
    dict(id="R1", kind="nyata", name="Layout 1 bulan ke-4: judul 4 baris dan kolom Catatan/Paraf", file="laporan_G1_2026-06.xlsx", mutate=None, old_without=["L1_v2"], new_cfg=None),
    dict(id="R2", kind="nyata", name="Layout 3 bulan ke-4: nama sheet 'Minggu 1' menjadi 'Wk1'", file="laporan_G3_2026-06.xlsx", mutate=None, old_without=["L3_v2"], new_cfg=None),
    dict(id="S1", kind="sintetis", name="Layout 4: satu kolom kosong disisipkan di depan", file="laporan_G4_2026-03.xlsx", mutate=m_insert_col_front("Output"), cfg="L4_v1",
         patches=[("- {cell: A1, equals: Shift 1}", "- {cell: B1, equals: Shift 1}"), ("- {cell: H1, equals: Shift 2}", "- {cell: I1, equals: Shift 2}"),
                  ("- {start_col: A, shift_cell: A1}", "- {start_col: B, shift_cell: B1}"), ("- {start_col: H, shift_cell: H1}", "- {start_col: I, shift_cell: I1}")]),
    dict(id="S2", kind="sintetis", name="Layout 2: dua baris judul ditambahkan di atas", file="laporan_G2_2026-03.xlsx", mutate=m_title_rows, cfg="L2_v1",
         patches=[("- {cell: A1, startswith: 'Tanggal:'}", "- {cell: A3, startswith: 'Tanggal:'}")]),
    dict(id="S3", kind="sintetis", name="Layout 5: kolom Operator disisipkan di tengah tabel", file="laporan_G5_2026-03.xlsx", mutate=m_insert_col_d, cfg="L5_v1",
         patches=[("columns: [date, shift, time, station, size, produced, reject]", "columns: [date, shift, time, _, station, size, produced, reject]"),
                  ("- {cell: F1, equals: Hasil}", "- {cell: G1, equals: Hasil}")]),
    dict(id="S4", kind="sintetis", name="Layout 1: sheet diganti nama dari 'Sheet1' menjadi 'Output'", file="laporan_G1_2026-03.xlsx", mutate=m_rename_sheet, cfg="L1_v1",
         patches=[("sheet_regex: '^Sheet1$'", "sheet_regex: '^Output$'"), ("sheets: {include: '^Sheet1$'}", "sheets: {include: '^Output$'}")]),
    dict(id="S5", kind="sintetis", name="Layout 6: sheet baru 'Catatan' ditambahkan di file", file="laporan_G6_2026-03.xlsx", mutate=m_extra_sheet, cfg="L6_v1", patches=[]),
    dict(id="S6", kind="sintetis", name="Layout 3: nama stasiun baru yang belum dikenal ('Filler 2')", file="laporan_G3_2026-03.xlsx", mutate=m_unknown_station, cfg="L3_v1", patches=[],
         lookup_patches=[("Filling: [Filling, Filler]", "Filling: [Filling, Filler, Filler 2]")]),
    dict(id="S7", kind="sintetis", name="Workbook yang bukan laporan produksi", file="laporan_G1_2026-03.xlsx", mutate=m_random, cfg=None, patches=[]),
]


def _patched_text(path, patches):
    t = Path(path).read_text(encoding="utf-8")
    for old, new in patches:
        assert t.count(old) == 1, (path, old)
        t = t.replace(old, new)
    return t


def _truth_index():
    t = load_truth_long()
    return t.set_index(KEY).truth


def score(fact, file, truth):
    """Cakupan dan ketepatan hasil parsing satu file terhadap kunci jawaban (hanya sel berstatus ok atau fixed)."""
    f = fact[fact.source_file == file].assign(date=lambda x: x.date.astype(str))
    tr = truth.reindex(pd.MultiIndex.from_frame(f[KEY]))
    good = f.quality.isin(["ok", "fixed"]).to_numpy()
    eq = (f.value.to_numpy() == tr.to_numpy()) & good
    return {"cells": len(f), "trusted": int(good.sum()), "trusted_exact": int(eq.sum())}


def run_scenario(s, tmp, truth):
    lk, rules = load_lookup(C.LOOKUP), load_rules(C.RULES)
    all_t = {c["id"]: c for c in load_templates(C.TEMPLATES_DIR)}
    src = C.REPORTS_DIR / s["file"]
    path = src
    if s.get("mutate"):
        wb = openpyxl.load_workbook(src)
        s["mutate"](wb)
        path = Path(tmp) / s["file"]
        wb.save(path)
    old_ids = [k for k in all_t if k not in s.get("old_without", [])]
    out = {"id": s["id"], "kind": s["kind"], "scenario": s["name"], "file": s["file"]}

    def attempt(templates, lookup):
        try:
            res = parse_files([path], templates, lookup, rules)
            return "ok", score(res["fact"], s["file"], truth), res
        except ParseError as e:
            return type(e).__name__, str(e), None

    st, info, res = attempt([all_t[i] for i in old_ids], lk)
    out["old_config"] = "berhenti: " + st if st != "ok" else "diparse"
    out["old_config_message"] = info if st != "ok" else ""
    if st == "ok":
        out["old_config"] = f"diparse, {info['trusted_exact']} dari {info['trusted']} sel tepercaya cocok"
        out["old_wrong"] = info["trusted"] - info["trusted_exact"]
    keys_changed = lookup_changed = 0
    templates_new, lookup_new = [all_t[i] for i in all_t], lk
    if s.get("old_without"):                                   # perubahan nyata: config baru sudah ada di repo
        a, b = s["old_without"][0], s["old_without"][0][:2] + "_v1"
        keys_changed = len(set(_lines(Path(C.TEMPLATES_DIR / f"{a}.yaml").read_text(encoding="utf-8"))) ^ set(_lines(Path(C.TEMPLATES_DIR / f"{b}.yaml").read_text(encoding="utf-8")))) // 2
    elif s.get("cfg") and (s["patches"]):
        txt = _patched_text(C.TEMPLATES_DIR / f"{s['cfg']}.yaml", s["patches"])
        patched = yaml.safe_load(txt)
        templates_new = [patched if c["id"] == s["cfg"] else c for c in all_t.values()]
        keys_changed = len(s["patches"])
    if s.get("lookup_patches"):
        d = yaml.safe_load(_patched_text(C.LOOKUP, s["lookup_patches"]))
        lookup_new = Lookup(d)
        lookup_changed = len(s["lookup_patches"])
    st2, info2, res2 = attempt(templates_new, lookup_new)
    out["new_config"] = "berhenti: " + st2 if st2 != "ok" else f"{info2['trusted_exact']} dari {info2['trusted']} sel tepercaya cocok"
    out["new_cells"] = info2["cells"] if st2 == "ok" else 0
    out["new_trusted"] = info2["trusted"] if st2 == "ok" else 0
    out["new_trusted_exact"] = info2["trusted_exact"] if st2 == "ok" else 0
    out["config_lines_changed"] = keys_changed
    out["lookup_lines_changed"] = lookup_changed
    out["python_lines_changed"] = 0
    out["old_stopped_safely"] = st != "ok" or s["id"] == "S5"
    return out


def run_all():
    truth = _truth_index()
    rows = []
    with tempfile.TemporaryDirectory() as tmp:
        for s in SCEN:
            rows.append(run_scenario(s, tmp, truth))
    return pd.DataFrame(rows)
