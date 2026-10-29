"""Empat metode untuk laporan layout 1 yang sama (Q8-Q10): rumus Excel (dihitung LibreOffice), SQL (DuckDB), Python (parser), Power Query (teks M)."""
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analysis_lib import C, KEY, load_truth_long  # noqa: E402
from parser import load_lookup, load_rules, load_templates, parse_files  # noqa: E402

N_ROWS, N_DATES = 36, 25            # baris per blok dan jumlah tanggal per file layout 1 (diketik manual di sheet Pemetaan)
FORMULA_COLS = {
    "A": ("k", None), "B": ("blok", '=IF(A{r}<={nr}*{nd},"PRODUKSI","REJECT")'), "C": ("urut", "=MOD(A{r}-1,Pemetaan!$B$3*Pemetaan!$B$4)"),
    "D": ("baris_ke", "=INT(C{r}/Pemetaan!$B$4)"), "E": ("kolom_ke", "=MOD(C{r},Pemetaan!$B$4)"),
    "F": ("baris_penanda", '=IF(B{r}="PRODUKSI",Pemetaan!$B$5,Pemetaan!$B$6)'), "G": ("baris_sumber", "=F{r}+2+D{r}"), "H": ("kolom_sumber", "=Pemetaan!$B$2+E{r}"),
    "I": ("station", "=INDEX(Laporan!$A$1:$AZ$300,G{r},1)"), "J": ("size", "=INDEX(Laporan!$A$1:$AZ$300,G{r},2)"),
    "K": ("shift", "=INDEX(Laporan!$A$1:$AZ$300,G{r},3)"), "L": ("slot", "=INDEX(Laporan!$A$1:$AZ$300,G{r},4)"),
    "M": ("date", "=DATE(Pemetaan!$B$1,VALUE(RIGHT(INDEX(Laporan!$A$1:$AZ$300,F{r}+1,H{r}),2)),VALUE(LEFT(INDEX(Laporan!$A$1:$AZ$300,F{r}+1,H{r}),2)))"),
    "N": ("mentah", "=INDEX(Laporan!$A$1:$AZ$300,G{r},H{r})"),
    "O": ("value", '=IF(N{r}="","",IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(N{r}," btl",""),".","")),""))'), "P": ("metric", '=IF(B{r}="PRODUKSI","produced","reject")'),
}


def lo_recalc(xlsx, outdir):
    """Hitung ulang rumus dengan LibreOffice (headless); mengembalikan jalur file hasil."""
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx)], check=True, capture_output=True, timeout=300)
    return Path(outdir) / Path(xlsx).name


def build_excel(src, out, first_value_col, year):
    """Buat workbook metode rumus: sheet Laporan (salinan isi), Pemetaan (parameter), Tidy (rumus INDEX)."""
    sw = openpyxl.load_workbook(src)["Sheet1"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan"
    for row in sw.iter_rows():
        for c in row:
            if c.value is not None:
                ws.cell(c.row, c.column, c.value)
    pm = wb.create_sheet("Pemetaan")
    for r, (label, val) in enumerate([("Tahun laporan", year), ("Kolom pertama nilai (nomor kolom)", first_value_col), ("Baris data per blok", N_ROWS), ("Jumlah tanggal", N_DATES),
                                      ("Baris penanda PRODUKSI", '=MATCH("PRODUKSI",Laporan!$A:$A,0)'), ("Baris penanda REJECT", '=MATCH("REJECT",Laporan!$A:$A,0)')], 1):
        pm.cell(r, 1, label)
        pm.cell(r, 2, val)
    td = wb.create_sheet("Tidy")
    for col, (name, _) in FORMULA_COLS.items():
        td[f"{col}1"] = name
    n = 2 * N_ROWS * N_DATES
    for i in range(1, n + 1):
        r = i + 1
        td[f"A{r}"] = i
        for col, (_, f) in FORMULA_COLS.items():
            if f:
                td[f"{col}{r}"] = f.format(r=r, nr=N_ROWS, nd=N_DATES)
    wb.save(out)
    return n


def run_excel(src, first_value_col, year, keep=None):
    """Bangun workbook, hitung ulang di LibreOffice, baca hasilnya. Mengembalikan (DataFrame, detik hitung ulang)."""
    with tempfile.TemporaryDirectory() as tmp:
        x = Path(tmp) / "m.xlsx"
        build_excel(src, x, first_value_col, year)
        t0 = time.perf_counter()
        y = lo_recalc(x, Path(tmp) / "out")
        sec = time.perf_counter() - t0
        if keep:
            Path(keep).write_bytes(y.read_bytes())
        ws = openpyxl.load_workbook(y, data_only=True)["Tidy"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, min_col=9, max_col=16)]
    df = pd.DataFrame(rows, columns=["station", "size", "shift", "slot", "date", "mentah", "value", "metric"]).drop(columns="mentah")
    df["date"] = pd.to_datetime(df["date"], errors="coerce", format="%Y-%m-%d %H:%M:%S").dt.strftime("%Y-%m-%d")
    df["slot"] = pd.to_numeric(df["slot"], errors="coerce")
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    return df, sec


def sql_text():
    t = (C.SQL_ANALYSIS / "method_L1.sql").read_text(encoding="utf-8")
    return t.split("-- name: method_L1", 1)[1].strip()


def run_sql(src, first_value_col, year):
    """Isi tabel cell_dump dari sheet laporan, jalankan query. Mengembalikan (DataFrame, detik total)."""
    t0 = time.perf_counter()
    ws = openpyxl.load_workbook(src)["Sheet1"]
    dump = [(c.row, c.column, str(c.value)) for row in ws.iter_rows() for c in row if c.value is not None]
    con = duckdb.connect()
    con.execute("CREATE TABLE cell_dump (r INTEGER, c INTEGER, v VARCHAR)")
    con.executemany("INSERT INTO cell_dump VALUES (?, ?, ?)", dump)
    df = con.execute(sql_text().format(first_value_col=first_value_col, year=year)).df()
    con.close()
    sec = time.perf_counter() - t0
    df["date"] = pd.to_datetime(df["date"]).dt.strftime("%Y-%m-%d")
    return df, sec


def run_python(paths):
    t0 = time.perf_counter()
    try:
        out = parse_files(paths, load_templates(C.TEMPLATES_DIR), load_lookup(C.LOOKUP), load_rules(C.RULES))
    except Exception as e:                                    # berhenti dengan pesan jelas (mis. template tidak dikenal)
        return None, time.perf_counter() - t0, f"{type(e).__name__}"
    f = out["fact"]
    f = f.assign(date=f.date.astype(str))
    return f[KEY + ["value"]], time.perf_counter() - t0, ""


def expected(files):
    """Sel kunci jawaban yang seharusnya ada pada file-file ini (kunci diambil dari hasil parsing, yang cakupannya sudah dicek 100%)."""
    f = pd.read_csv(C.PARSED_DIR / "fact_output_tidy.csv", usecols=KEY + ["source_file"])
    f = f[f.source_file.isin(files)]
    return f.merge(load_truth_long(), on=KEY, how="left", validate="one_to_one")


def match_truth(df, files):
    """Persen sel kunci jawaban pada file-file ini yang nilainya sama persis dengan hasil metode."""
    e = expected(files)
    if df is None or not len(df):
        return 0.0
    d = df.drop_duplicates(KEY)[KEY + ["value"]].astype({"value": "float64"})
    m = e.merge(d, on=KEY, how="left")
    return float((m.value == m.truth).mean())
