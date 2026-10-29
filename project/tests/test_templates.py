import copy

import openpyxl
import pytest

from parser import AmbiguousTemplate, UnknownTemplate, load_templates
from parser.config import detect_template


def test_configs_load_and_have_unique_ids(templates):
    ids = [t["id"] for t in templates]
    assert len(ids) == len(set(ids)) == 8
    assert sorted({t["layout"] for t in templates}) == [1, 2, 3, 4, 5, 6]


def test_every_report_matches_exactly_one_template(sim, templates):
    for p in sorted(sim.glob("*.xlsx")):
        t = detect_template(openpyxl.load_workbook(p), templates)
        assert t["layout"] == int(p.name.split("_G")[1][0]), p.name


def test_layout_change_picks_v2(sim, templates):
    pick = lambda n: detect_template(openpyxl.load_workbook(sim / n), templates)["id"]  # noqa: E731
    assert pick("laporan_G1_2026-05.xlsx") == "L1_v1" and pick("laporan_G1_2026-06.xlsx") == "L1_v2"
    assert pick("laporan_G3_2026-05.xlsx") == "L3_v1" and pick("laporan_G3_2026-06.xlsx") == "L3_v2"


def test_unknown_workbook_is_refused(templates):
    wb = openpyxl.Workbook()
    wb.active.title = "Hitungan"
    with pytest.raises(UnknownTemplate):
        detect_template(wb, templates)


def test_ambiguous_templates_are_refused(sim, templates):
    dup = copy.deepcopy(templates[0])
    dup["id"] = "dup"
    wb = openpyxl.load_workbook(sim / "laporan_G1_2026-03.xlsx")
    base = [t for t in templates if t["id"] == "L1_v1"][0]
    dup = copy.deepcopy(base)
    dup["id"] = "L1_copy"
    with pytest.raises(AmbiguousTemplate):
        detect_template(wb, [base, dup])
